"""Sales-workbook parser + DB inserter.

The founder maintains a daily Excel file per day in folders like
`<root>/<MonthGE> NN YYYY/DD.MM.YYYY.xlsx`. Two formats coexist:

  - English layout (Apr-Jun 2026 onward): sheet "Sales", header at row 3:
    CODE / QUANTITY / UNIT PRICE / TOTAL / COMPANY.
  - Turkish legacy (Jan-Mar 2026): sheet "satis", headers in Turkish but
    column positions identical. Same date format in C2 (`DD/MM/YYYY`).

This module exposes parsing as pure functions so both the one-off CLI
(`scripts/import_sales_2026.py`) and the Telegram `/import_sales` handler
can reuse it.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator, Optional

import openpyxl

from src import db
from src.logging_setup import get_logger

log = get_logger(__name__)

# Column indexes (0-based) inside the Sales / satis sheet.
COL_CODE = 2
COL_QTY = 3
COL_UNIT_PRICE = 4
COL_TOTAL = 5
COL_COMPANY = 6

SHEET_NAMES = ("Sales", "satis")


@dataclass
class ParsedSale:
    product_code: str
    quantity: float
    unit_price: float
    total_price: float
    sold_at: datetime
    company: Optional[str] = None
    source_file: Optional[str] = None


# ─── Coercion helpers ────────────────────────────────────────────────────────


def _normalize_code(raw: object) -> Optional[str]:
    if raw is None:
        return None
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip() or None


def _coerce_float(raw: object) -> Optional[float]:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    try:
        return float(str(raw).replace(",", "."))
    except ValueError:
        return None


def _parse_date(raw: object) -> Optional[datetime]:
    if isinstance(raw, datetime):
        return raw
    if not isinstance(raw, str):
        return None
    s = raw.strip()
    for fmt in ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


# ─── Workbook parsing ────────────────────────────────────────────────────────


def parse_workbook(path: Path) -> tuple[Optional[datetime], list[ParsedSale]]:
    """Open one .xlsx file. Return (sale_date, parsed_rows).

    Returns (None, []) if the workbook has no recognized sales sheet — caller
    can skip the file.
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        log.warning("xlsx_open_failed", path=str(path))
        return None, []

    sheet_name = next((n for n in SHEET_NAMES if n in wb.sheetnames), None)
    if sheet_name is None:
        return None, []

    ws = wb[sheet_name]
    sold_at: Optional[datetime] = None
    rows: list[ParsedSale] = []

    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_idx == 2:
            sold_at = _parse_date(row[COL_CODE] if len(row) > COL_CODE else None)
            continue
        if r_idx < 4:
            continue
        if len(row) <= COL_TOTAL:
            continue
        code = _normalize_code(row[COL_CODE])
        qty = _coerce_float(row[COL_QTY])
        if code is None or qty is None or qty == 0:
            continue
        unit_price = _coerce_float(row[COL_UNIT_PRICE])
        total = _coerce_float(row[COL_TOTAL])
        if total is None and unit_price is not None:
            total = round(qty * unit_price, 2)
        if unit_price is None and total is not None and qty:
            unit_price = round(total / qty, 4)
        if total is None or unit_price is None:
            continue
        company = row[COL_COMPANY] if len(row) > COL_COMPANY else None
        company = str(company).strip() if company else None
        rows.append(ParsedSale(
            product_code=code,
            quantity=qty,
            unit_price=unit_price,
            total_price=total,
            # sold_at filled below after we know the date
            sold_at=datetime.min,
            company=company,
            source_file=path.name,
        ))

    if sold_at is None:
        # Fall back to DD.MM.YYYY filename.
        m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", path.name)
        if m:
            sold_at = datetime.strptime(".".join(m.groups()), "%d.%m.%Y")

    if sold_at is None:
        return None, []

    for r in rows:
        r.sold_at = sold_at
    return sold_at, rows


def parse_directory(root: Path) -> Iterator[ParsedSale]:
    """Yield every parsed sale from .xlsx files under `root` (recursive)."""
    for path in sorted(root.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue  # Excel lock file
        _, rows = parse_workbook(path)
        yield from rows


def parse_zip(zip_path: Path, extract_to: Path) -> Iterator[ParsedSale]:
    """Unzip into `extract_to`, then parse every .xlsx found inside."""
    extract_to.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as zf:
        zf.extractall(extract_to)
    yield from parse_directory(extract_to)


def parse_files(paths: Iterable[Path]) -> Iterator[ParsedSale]:
    """Parse an explicit list of .xlsx file paths."""
    for path in paths:
        if path.name.startswith("~$") or path.suffix.lower() != ".xlsx":
            continue
        _, rows = parse_workbook(path)
        yield from rows


# ─── DB persistence ──────────────────────────────────────────────────────────


def commit_sales(
    parsed: list[ParsedSale],
    *,
    wipe_period: Optional[tuple[datetime, datetime]] = None,
) -> tuple[int, int]:
    """Insert parsed sales into `sales` table.

    `wipe_period=(start, end)`: delete existing sales in [start, end) before
    insert. Used by the monthly importer for idempotent re-runs.

    Returns (inserted_count, deleted_count).
    """
    if not parsed:
        return 0, 0

    deleted = 0
    if wipe_period is not None:
        start, end = wipe_period
        with db.session_scope() as s:
            deleted = (
                s.query(db.Sale)
                .filter(db.Sale.sold_at >= start, db.Sale.sold_at < end)
                .delete(synchronize_session=False)
            )

    with db.session_scope() as s:
        catalog = {p.code: p.name for p in s.query(db.Product).all()}
        objs = [
            db.Sale(
                product_code=r.product_code,
                product_name=catalog.get(r.product_code, "?"),
                quantity=r.quantity,
                unit_price=r.unit_price,
                total_price=r.total_price,
                notes=r.company,
                sold_at=r.sold_at,
            )
            for r in parsed
        ]
        s.bulk_save_objects(objs)
    return len(parsed), deleted


# ─── Helpers for callers ─────────────────────────────────────────────────────


def summarize(parsed: list[ParsedSale]) -> dict:
    """Aggregate stats for a quick UX preview before commit."""
    if not parsed:
        return {"row_count": 0, "revenue": 0.0, "files": 0, "date_min": None,
                "date_max": None, "unique_codes": 0, "unknown_codes": []}

    revenue = sum(r.total_price for r in parsed)
    dates = sorted({r.sold_at.date() for r in parsed})
    codes = {r.product_code for r in parsed}
    files = {r.source_file for r in parsed if r.source_file}

    with db.session_scope() as s:
        known = {p.code for p in s.query(db.Product).all()}
    unknown = sorted(codes - known)

    return {
        "row_count": len(parsed),
        "revenue": round(revenue, 2),
        "files": len(files),
        "date_min": dates[0].isoformat(),
        "date_max": dates[-1].isoformat(),
        "unique_codes": len(codes),
        "unknown_codes": unknown,
    }


def month_bounds(year: int, month: int) -> tuple[datetime, datetime]:
    """Return [start_of_month, start_of_next_month) as datetime pair."""
    start = datetime(year, month, 1)
    if month == 12:
        end = datetime(year + 1, 1, 1)
    else:
        end = datetime(year, month + 1, 1)
    return start, end


def filter_by_period(
    parsed: Iterable[ParsedSale], year: int, month: int
) -> list[ParsedSale]:
    """Keep only parsed rows whose sold_at falls inside the given month."""
    start, end = month_bounds(year, month)
    return [r for r in parsed if start <= r.sold_at < end]
