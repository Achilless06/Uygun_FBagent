"""Import products.xlsx into SQLite.

Source: `products.xlsx`, sheet `პროდუქცია`, columns: კოდი / დასახელება / ფასი.
481 product rows + 1 header. Codes are heterogeneous (mostly numeric strings,
but some are Georgian letters like "ც", "7ქ").

This module is both a library (`run_import`) and a CLI script:

    python -m src.catalog.importer            # uses default path products.xlsx
    python -m src.catalog.importer some.xlsx  # specify path

Strategy: **additive upsert**.
  - Existing products: update name/price/category if changed.
  - New products: insert with stock_qty=1.
  - Products in DB but not in the new file: leave alone (do NOT delete).
    Rationale: a bad import shouldn't wipe inventory; founder removes via
    Telegram if needed.

Data quality observed during profiling:
  - 14 rows have null or zero price → stored as NULL (agent treats as
    "ask for price" mode, won't include in price-mention posts).
  - 2 rows have null name → skipped with a warning.

Category derivation: first word of the product name. Crude but useful for
the generator's "pick a varied set of products this week" logic.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from src import config, db
from src.logging_setup import get_logger

log = get_logger(__name__)

DEFAULT_XLSX_PATH = config.PROJECT_ROOT / "products.xlsx"
SHEET_NAME = "პროდუქცია"

# Photos folder is scanned during import to set has_photo flag.
PHOTOS_DIR = config.PHOTOS_DIR


@dataclass
class ImportSummary:
    inserted: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped: int = 0

    def __str__(self) -> str:
        return (
            f"inserted={self.inserted} updated={self.updated} "
            f"unchanged={self.unchanged} skipped={self.skipped}"
        )


def _derive_category(name: str) -> str | None:
    """First word of the product name (e.g. 'საბურავის', 'რეზინის', 'AERO').

    Used by the generator to pick a varied mix of categories across the week.
    """
    if not name:
        return None
    return name.strip().split()[0] if name.strip() else None


def _normalize_price(raw) -> float | None:
    """Coerce price cell to float or None. Zero/negative → None."""
    if raw is None:
        return None
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _natural_sort_key(code: str) -> int | None:
    """Return int(code) if the code is a pure integer, None otherwise.

    Used to make `/products` sort 1, 2, ..., 10, 11, ... instead of the
    text-sort order 1, 10, 100, 11, 2, ...
    Non-numeric codes (e.g. "ც", "7ქ") get NULL and sort to the end.
    """
    try:
        return int(code)
    except (ValueError, TypeError):
        return None


def _scan_photos() -> set[str]:
    """Return product codes that have a photo file in data/photos/.

    Looks for files named {code}.jpg, .jpeg, or .png.
    """
    if not PHOTOS_DIR.exists():
        return set()
    codes = set()
    for path in PHOTOS_DIR.iterdir():
        if path.suffix.lower() in (".jpg", ".jpeg", ".png"):
            codes.add(path.stem)
    return codes


def run_import(xlsx_path: Path | None = None) -> ImportSummary:
    """Import products from the given xlsx file. Returns counts."""
    path = xlsx_path or DEFAULT_XLSX_PATH
    if not path.exists():
        raise FileNotFoundError(f"products.xlsx not found at {path}")

    log.info("import_start", path=str(path))

    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' not found in {path}. Sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    photos_with_files = _scan_photos()
    summary = ImportSummary()

    with db.session_scope() as session:
        existing = {p.code: p for p in session.query(db.Product).all()}

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            # Skip header row.
            if row_idx == 1:
                continue

            raw_code, raw_name, raw_price = (row + (None, None, None))[:3]

            if raw_code is None or raw_name is None:
                summary.skipped += 1
                log.warning("import_skip_row", row=row_idx, reason="null_code_or_name")
                continue

            code = str(raw_code).strip()
            name = str(raw_name).strip()
            if not code or not name:
                summary.skipped += 1
                log.warning("import_skip_row", row=row_idx, reason="empty_after_strip")
                continue

            price = _normalize_price(raw_price)
            category = _derive_category(name)
            has_photo = code in photos_with_files
            code_sort = _natural_sort_key(code)

            current = existing.get(code)
            if current is None:
                # New product.
                session.add(
                    db.Product(
                        code=code,
                        code_sort=code_sort,
                        name=name,
                        price=price,
                        stock_qty=1,
                        category=category,
                        has_photo=has_photo,
                    )
                )
                summary.inserted += 1
            else:
                # Detect changes; skip update if identical.
                changed = (
                    current.name != name
                    or current.price != price
                    or current.category != category
                    or current.has_photo != has_photo
                    or current.code_sort != code_sort
                )
                if changed:
                    current.name = name
                    current.price = price
                    current.category = category
                    current.has_photo = has_photo
                    current.code_sort = code_sort
                    summary.updated += 1
                else:
                    summary.unchanged += 1

    log.info("import_done", **summary.__dict__)
    return summary


def _main(argv: list[str]) -> int:
    """CLI entrypoint."""
    from src.logging_setup import configure as configure_logging

    configure_logging("INFO")

    # Initialize DB without going through config.load() — that requires every env
    # var to be set, which is overkill just to run the importer locally.
    db_url = f"sqlite:///{config.DATA_DIR}/uygun.db"
    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    db.init_engine(db_url)
    db.create_all()
    db.seed_default_settings()

    xlsx_path = Path(argv[1]) if len(argv) > 1 else DEFAULT_XLSX_PATH
    summary = run_import(xlsx_path)
    print(f"Import complete: {summary}")

    # Quick stats for the smoke test.
    with db.session_scope() as session:
        total = session.query(db.Product).count()
        with_photos = session.query(db.Product).filter(db.Product.has_photo.is_(True)).count()
        priced = session.query(db.Product).filter(db.Product.price.isnot(None)).count()
        print(f"DB now contains: {total} products ({priced} priced, {with_photos} with photos)")

    return 0


if __name__ == "__main__":
    sys.exit(_main(sys.argv))
